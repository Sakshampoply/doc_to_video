import os
import re
import time
import logging
import tempfile
from dotenv import load_dotenv
from moviepy import ImageClip, concatenate_videoclips, AudioFileClip
from PIL import Image, ImageDraw, ImageFont
import PyPDF2
import docx
from openai import AzureOpenAI
from gtts import gTTS

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Load environment variables
load_dotenv()

class Config:
    """Application configuration settings"""
    SLIDE_SIZE = (1920, 1080)
    FONT_PATH = "arial.ttf"
    TITLE_FONT_SIZE = 56
    CONTENT_FONT_SIZE = 36
    MAX_SUMMARY_RETRIES = 3
    SUMMARY_TIMEOUT = 30
    MIN_SLIDE_LENGTH = 150
    MAX_SLIDE_LENGTH = 600
    VIDEO_FPS = 24
    CONTENT_CLEAN_PATTERNS = [
        (r'\d+\.\s*', ''),         # Remove numbered bullets
        (r'[#*\-]{2,}', ''),       # Remove markdown symbols
        (r'\bSlide\s\d+\b', ''),   # Remove slide numbers
        (r'\s+', ' ')              # Consolidate whitespace
    ]
    NARRATION_CLEAN_PATTERNS = [
        (r'[^a-zA-Z0-9\s.,:;\-]', ''),  # Remove special chars
        (r'\s+', ' ')                   # Normalize whitespace
    ]

class DocumentProcessor:
    """Handles document processing and content generation"""
    def __init__(self):
        self.azure_client = AzureOpenAI(
            api_key=os.getenv("AZURE_OPENAI_API_KEY"),
            azure_endpoint=os.getenv("AZURE_OPENAI_ENDPOINT"),
            api_version="2024-02-01-preview",
            timeout=Config.SUMMARY_TIMEOUT
        )

    def process_document(self, file_path):
        """Main processing pipeline"""
        text = self._extract_text(file_path)
        if not text:
            raise ValueError("Failed to extract text from document")
            
        slides = self._generate_slides(text)
        if not slides:
            raise ValueError("Failed to generate meaningful slides")
            
        return slides

    def _extract_text(self, file_path):
        """Extract text from supported file types"""
        try:
            if file_path.endswith(".pdf"):
                return self._extract_pdf(file_path)
            if file_path.endswith(".docx"):
                return self._extract_docx(file_path)
            raise ValueError("Unsupported file format")
        except Exception as e:
            logger.error(f"Text extraction failed: {e}")
            return ""

    def _extract_pdf(self, pdf_path):
        """PDF text extraction with error handling"""
        text = []
        try:
            with open(pdf_path, "rb") as file:
                reader = PyPDF2.PdfReader(file)
                for page in reader.pages:
                    if page_text := page.extract_text():
                        text.append(page_text)
            return "\n".join(text)
        except Exception as e:
            logger.error(f"PDF extraction error: {e}")
            return ""

    def _extract_docx(self, docx_path):
        """DOCX text extraction with error handling"""
        try:
            doc = docx.Document(docx_path)
            return "\n".join(para.text for para in doc.paragraphs if para.text)
        except Exception as e:
            logger.error(f"DOCX extraction error: {e}")
            return ""

    def _generate_slides(self, text):
        """Generate and clean presentation slides"""
        raw_slides = self._azure_summarization(text)
        return self._clean_content(raw_slides)

    def _azure_summarization(self, text):
        """Generate slides using Azure OpenAI with enhanced prompts"""
        summarization_prompt = """Create a professional presentation from this document:
        - Use natural language without markdown
        - Format each slide as "Title: Clear content"
        - Focus on key concepts and relationships
        - Omit slide numbers and section markers
        - Maintain technical accuracy
        - Use concise bullet points when appropriate"""
        
        for attempt in range(Config.MAX_SUMMARY_RETRIES):
            try:
                response = self.azure_client.chat.completions.create(
                    model="anthropic.claude-v3-sonnet",
                    messages=[
                        {"role": "system", "content": summarization_prompt},
                        {"role": "user", "content": f"Document content:\n{text[:30000]}"}
                    ],
                    temperature=0.2,
                    max_tokens=2000
                )
                return self._split_slides(response.choices[0].message.content)
            except Exception as e:
                logger.warning(f"Summarization attempt {attempt+1} failed: {e}")
                time.sleep(2 ** attempt)
        return []

    def _split_slides(self, content):
        """Split generated content into individual slides"""
        return [slide.strip() for slide in re.split(r'\n\s*\n+', content) 
                if Config.MIN_SLIDE_LENGTH <= len(slide) <= Config.MAX_SLIDE_LENGTH]

    def _clean_content(self, slides):
        """Remove unwanted formatting from slides"""
        cleaned = []
        for slide in slides:
            for pattern, replacement in Config.CONTENT_CLEAN_PATTERNS:
                slide = re.sub(pattern, replacement, slide)
            if slide.count(':') == 1:
                title, content = slide.split(':', 1)
                slide = f"{title.strip()}: {content.strip()}"
            cleaned.append(slide)
        return cleaned

class MediaGenerator:
    """Handles slide and audio generation"""
    def __init__(self):
        self.title_font = self._load_font(Config.TITLE_FONT_SIZE)
        self.content_font = self._load_font(Config.CONTENT_FONT_SIZE)

    def _load_font(self, size):
        """Load font with fallback"""
        try:
            return ImageFont.truetype(Config.FONT_PATH, size)
        except IOError:
            logger.warning("Using default system font")
            return ImageFont.load_default(size)

    def create_slide(self, text, output_path):
        """Generate visual slide from text content"""
        img = Image.new("RGB", Config.SLIDE_SIZE, (30, 30, 50))
        draw = ImageDraw.Draw(img)
        
        title, content = self._split_title_content(text)
        self._draw_title(draw, title)
        self._draw_content(draw, content)
        
        img.save(output_path)

    def _split_title_content(self, text):
        """Separate title and body content"""
        if ':' in text:
            title, content = text.split(':', 1)
            return title.strip(), content.strip()
        return "Key Point", text

    def _draw_title(self, draw, title):
        """Render slide title"""
        bbox = draw.textbbox((0, 0), title, font=self.title_font)
        x = (Config.SLIDE_SIZE[0] - bbox[2]) // 2
        draw.text((x, 100), title, font=self.title_font, fill="white")

    def _draw_content(self, draw, content):
        """Render slide body content"""
        y = 200
        for line in self._wrap_text(content):
            draw.text((100, y), line, font=self.content_font, fill="white")
            y += self.content_font.size + 10

    def _wrap_text(self, text):
        """Wrap text to fit slide width"""
        lines = []
        max_width = Config.SLIDE_SIZE[0] - 200
        paragraphs = re.split(r'\.\s+', text)
        
        for para in paragraphs:
            words = para.split()
            current_line = []
            
            for word in words:
                test_line = ' '.join(current_line + [word])
                if self.content_font.getlength(test_line) < max_width:
                    current_line.append(word)
                else:
                    lines.append(' '.join(current_line))
                    current_line = [word]
            if current_line:
                lines.append(' '.join(current_line))
        
        return lines

    def generate_narration(self, text, output_path):
        """Generate audio narration from text"""
        clean_text = self._clean_narration_text(text)
        try:
            tts = gTTS(text=clean_text, lang='en', slow=False)
            tts.save(output_path)
            return AudioFileClip(output_path).duration
        except Exception as e:
            logger.error(f"Narration failed: {e}")
            return 5  # Fallback duration

    def _clean_narration_text(self, text):
        """Prepare text for TTS conversion"""
        clean_text = text.split(':', 1)[-1]  # Remove title
        for pattern, replacement in Config.NARRATION_CLEAN_PATTERNS:
            clean_text = re.sub(pattern, replacement, clean_text)
        return clean_text.strip()

class VideoPipeline:
    """Manages video creation workflow"""
    def __init__(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.clips = []
        self.audio_clips = []

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        """Ensure proper resource cleanup"""
        for clip in self.clips + self.audio_clips:
            try:
                clip.close()
            except AttributeError:
                pass
        
        time.sleep(1)  # Allow OS to release file locks
        try:
            self.temp_dir.cleanup()
        except PermissionError:
            logger.warning("Some temporary files could not be removed")

    def add_slide(self, slide_text, media_gen):
        """Create and add a slide to the pipeline"""
        idx = len(self.clips)
        img_path = os.path.join(self.temp_dir.name, f"slide_{idx}.png")
        audio_path = os.path.join(self.temp_dir.name, f"audio_{idx}.mp3")

        # Generate media
        media_gen.create_slide(slide_text, img_path)
        duration = media_gen.generate_narration(slide_text, audio_path)

        # Create and track clips
        img_clip = ImageClip(img_path).with_duration(duration)
        audio_clip = AudioFileClip(audio_path)
        self.audio_clips.append(audio_clip)
        
        self.clips.append(img_clip.with_audio(audio_clip))

    def render_video(self, output_path):
        """Final video rendering"""
        video = concatenate_videoclips(self.clips)
        video.write_videofile(output_path, fps=Config.VIDEO_FPS)
        video.close()
        logger.info(f"Successfully created video: {output_path}")

def main(input_file, output_file="output.mp4"):
    """Main execution flow"""
    try:
        processor = DocumentProcessor()
        media_gen = MediaGenerator()
        
        with VideoPipeline() as pipeline:
            slides = processor.process_document(input_file)
            logger.info(f"Generated {len(slides)} slides")
            
            for slide in slides:
                pipeline.add_slide(slide, media_gen)
                
            pipeline.render_video(output_file)
            
    except Exception as e:
        logger.error(f"Processing failed: {e}")
        raise

if __name__ == "__main__":
    input_path = r"C:\Users\saksham_poply\Downloads\Conversational AI for Media Search- A Retrieval-Augmented Generation Approach with LangChain (1).pdf"
    main(input_path)